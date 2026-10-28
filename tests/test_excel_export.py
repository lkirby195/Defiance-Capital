"""outputs/excel.py: the workbook a person checks the math in.  # SPEC §9.6

Two things are being asserted. First that every figure is written as a *number* under a
format - a workbook of preformatted strings adds up to nothing and is the failure mode this
export exists to avoid. Second, and the reason the Return Overview sheet is shaped the way it
is, that the ledger it writes down really does produce the IRR the engine claims: the test
reads the dates and the ``net`` column back out of the sheet, solves the XIRR itself with an
implementation that shares no code with ``engine/calc/irr.py``, and requires the two to agree
to 1e-6. The sheet carries Excel's own ``XIRR`` formula over the same two ranges, so a reader
opening the file sees the same check recomputed in front of them.
"""

from __future__ import annotations

import json
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from cli.fixtures import FixtureRun, run_fixture
from config.config import Config
from outputs.excel import build_workbook, write_workbook

FIXTURE_DIR = Path(__file__).resolve().parents[1] / "fixtures/synthetic/deals"
CONFIG = Config.load()
D = Decimal

SHEETS = [
    "Inputs",
    "Return Overview",
    "Lender Sensitivity",
    "Borrower Sensitivity",
    "Pricing",
    "Flip",
    "Rental",
    "Take-Back",
    "Sources & Uses",
    "Flags",
]


def underwritten() -> list[Path]:
    return [
        path
        for path in sorted(FIXTURE_DIR.glob("*.json"))
        if "underwrite" in json.loads(path.read_text(encoding="utf-8"))
    ]


UNDERWRITE_FILES = underwritten()


def run(path: Path) -> FixtureRun:
    return run_fixture(path, CONFIG, with_underwrite=True)


def book_for(path: Path) -> Any:
    fixture = run(path)
    assert fixture.underwrite_inputs is not None and fixture.underwrite_result is not None
    return build_workbook(
        fixture.screen_inputs, fixture.underwrite_inputs, fixture.underwrite_result, CONFIG
    )


def xirr_by_secant(flows: list[tuple[date, float]]) -> float:
    """An independent XIRR: plain floats, the secant method, no shared code with the engine.

    Deliberately not Newton and deliberately not Decimal - if this and the engine agree to a
    millionth, they are agreeing about the arithmetic rather than about an implementation.
    """
    start = flows[0][0]
    years = [(when - start).days / 365.0 for when, _ in flows]
    amounts = [amount for _, amount in flows]

    def npv(rate: float) -> float:
        return sum(a / (1.0 + rate) ** t for a, t in zip(amounts, years, strict=True))

    low, high = 0.0, 0.2
    f_low, f_high = npv(low), npv(high)
    for _ in range(200):
        if f_high == f_low:
            break
        nxt = high - f_high * (high - low) / (f_high - f_low)
        if nxt <= -0.999:
            nxt = (high - 0.999) / 2
        low, f_low, high = high, f_high, nxt
        f_high = npv(high)
        if abs(high - low) < 1e-14:
            break
    return high


def ledger_rows(sheet: Worksheet) -> list[tuple[date, float]]:
    """The dates and the net column, read back out of the sheet as a person would."""
    rows: list[tuple[date, float]] = []
    for row in sheet.iter_rows(min_row=5, max_col=9):
        when, net = row[0].value, row[8].value
        if not isinstance(when, datetime | date) or not isinstance(net, int | float | Decimal):
            break
        rows.append((when.date() if isinstance(when, datetime) else when, float(net)))
    return rows


def find_row(sheet: Worksheet, label: str) -> tuple[Any, Any]:
    for row in sheet.iter_rows(min_col=1, max_col=3):
        if row[0].value == label:
            return row[1].value, row[1].number_format
    raise AssertionError(f"no row labelled {label!r} on {sheet.title}")


# --- the sheets (SPEC §9.6) ----------------------------------------------------------------------


def test_the_workbook_has_its_sheets_in_the_spec_order() -> None:
    assert book_for(UNDERWRITE_FILES[0]).sheetnames == SHEETS


@pytest.mark.parametrize("path", UNDERWRITE_FILES, ids=[p.stem for p in UNDERWRITE_FILES])
def test_every_fixture_exports(path: Path) -> None:
    assert book_for(path).sheetnames == SHEETS


def test_write_workbook_saves_a_file_that_opens(tmp_path: Path) -> None:
    fixture = run(UNDERWRITE_FILES[0])
    assert fixture.underwrite_inputs is not None and fixture.underwrite_result is not None
    out = tmp_path / "deal.xlsx"
    written = write_workbook(
        out, fixture.screen_inputs, fixture.underwrite_inputs, fixture.underwrite_result, CONFIG
    )
    assert written == out and out.exists()
    assert load_workbook(out).sheetnames == SHEETS


# --- numbers as numbers --------------------------------------------------------------------------


@pytest.mark.parametrize("path", UNDERWRITE_FILES, ids=[p.stem for p in UNDERWRITE_FILES])
def test_the_ledger_is_written_as_dates_and_numbers(path: Path) -> None:
    sheet = book_for(path)["Return Overview"]
    result = run(path).underwrite_result
    assert result is not None
    at = 5
    for entry in result.return_overview.entries:
        when = sheet.cell(row=at, column=1)
        assert isinstance(when.value, datetime | date)
        assert when.number_format == "yyyy-mm-dd"
        assert sheet.cell(row=at, column=2).value == entry.month
        # The stub column is blank on a whole period and carries its days on the short one.
        assert sheet.cell(row=at, column=3).value == (entry.stub_days or None)
        for column, amount in enumerate(
            (entry.funding, entry.draws, entry.interest, entry.fees, entry.payoff, entry.net),
            start=4,
        ):
            cell = sheet.cell(row=at, column=column)
            assert isinstance(cell.value, Decimal | int | float), (at, column)
            assert Decimal(str(cell.value)) == amount
            assert cell.number_format == '"$"#,##0.00'
        at += 1


def test_percentages_are_percentages_and_ratios_are_ratios() -> None:
    book = book_for(UNDERWRITE_FILES[0])
    _, fmt = find_row(book["Inputs"], "Interest rate")
    assert fmt == "0.0%"
    _, fmt = find_row(book["Flip"], "Project yield")
    assert fmt == "0.0%"
    _, fmt = find_row(book["Flip"], "Cash-on-cash")
    assert fmt == "0.0%"
    _, fmt = find_row(book["Rental"], "Target DSCR")
    assert fmt == '0.00"x"'
    _, fmt = find_row(book["Rental"], "DSCR at market, amortizing")
    assert fmt == '0.00"x"'
    _, fmt = find_row(book["Inputs"], "Closing date")
    assert fmt == "yyyy-mm-dd"


def test_no_money_cell_is_a_string() -> None:
    """The failure this export exists to avoid: a workbook of text that adds up to nothing."""
    for path in UNDERWRITE_FILES:
        book = book_for(path)
        for name in SHEETS:
            for row in book[name].iter_rows():
                for cell in row:
                    if cell.number_format == '"$"#,##0.00' and cell.value is not None:
                        # Excel's own arithmetic is a formula, not text: the Sources & Uses
                        # check row (SPEC §8.11) is the one money formula
                        if str(cell.value).startswith("="):
                            assert name == "Sources & Uses", (name, cell.coordinate)
                            continue
                        assert isinstance(cell.value, Decimal | int | float), (
                            name,
                            cell.coordinate,
                        )


# --- the IRR the workbook recomputes for itself (SPEC §9.6) --------------------------------------


@pytest.mark.parametrize("path", UNDERWRITE_FILES, ids=[p.stem for p in UNDERWRITE_FILES])
def test_the_sheet_carries_an_xirr_formula_over_its_own_ledger(path: Path) -> None:
    sheet = book_for(path)["Return Overview"]
    result = run(path).underwrite_result
    assert result is not None
    last = 4 + len(result.return_overview.entries)
    formula, fmt = find_row(sheet, "IRR (this workbook)")
    assert formula == f"=XIRR(I5:I{last},A5:A{last})"
    assert fmt == "0.0000%"
    # ...over exactly the rows the ledger occupies, and no others
    assert len(ledger_rows(sheet)) == len(result.return_overview.entries)


@pytest.mark.parametrize("path", UNDERWRITE_FILES, ids=[p.stem for p in UNDERWRITE_FILES])
def test_the_workbook_ledger_solves_to_the_engines_irr(path: Path) -> None:
    """Read the sheet's own dates and net column back out and solve them independently."""
    sheet = book_for(path)["Return Overview"]
    result = run(path).underwrite_result
    assert result is not None and result.return_overview.irr is not None

    from_sheet = xirr_by_secant(ledger_rows(sheet))
    assert abs(from_sheet - float(result.return_overview.irr)) < 1e-6, path.stem

    engine_irr, fmt = find_row(sheet, "IRR (engine)")
    assert fmt == "0.0000%"
    assert abs(float(engine_irr) - from_sheet) < 1e-6


# --- the sensitivity sheet (SPEC §8.9) ----------------------------------------------------------


def test_the_sensitivity_sheet_is_numbers_under_formats_with_the_request_on_the_bottom_row() -> (
    None
):
    path = FIXTURE_DIR / "go_holdback_denver.json"
    result = run(path).underwrite_result
    assert result is not None and result.sensitivity is not None
    sheet = book_for(path)["Lender Sensitivity"]
    # the rate columns are numbers under a percent format, not text
    rates = [
        sheet.cell(row=4, column=column) for column in range(5, 5 + len(result.sensitivity.rates))
    ]
    assert [Decimal(str(cell.value)) for cell in rates] == result.sensitivity.rates
    assert all(cell.number_format == "0.0%" for cell in rates)
    # one row per loan amount, the request last, every IRR a number under a percent format
    rows = result.sensitivity.rows
    for offset, row in enumerate(rows):
        at = 5 + offset
        amount = sheet.cell(row=at, column=1)
        assert (
            Decimal(str(amount.value)) == row.loan_amount and amount.number_format == '"$"#,##0.00'
        )
        assert Decimal(str(sheet.cell(row=at, column=2).value)) == -row.reduction
        for column, cell in enumerate(row.cells, start=5):
            written = sheet.cell(row=at, column=column)
            assert cell.irr is not None and Decimal(str(written.value)) == cell.irr
            assert written.number_format == "0.0000%"
    assert (
        Decimal(str(sheet.cell(row=4 + len(rows), column=1).value))
        == result.economics.loan_requested
    )
    # the Denver deal is priced at 12%, off the grid: the columns start at 12.5% and nothing
    # is shaded, and the note under the title says why
    assert Decimal(str(rates[0].value)) == Decimal("0.125")
    shaded = [
        (offset, column)
        for offset in range(len(rows))
        for column in range(5, 5 + len(result.sensitivity.rates))
        if sheet.cell(row=5 + offset, column=column).fill.fgColor.rgb == "00E7F3EC"
    ]
    assert shaded == []
    assert "(12.0%), which is not a column" in str(sheet.cell(row=2, column=1).value)


def test_the_deals_own_cell_is_shaded_when_its_rate_is_a_column() -> None:
    from engine.underwrite import underwrite
    from tests.test_sensitivity import denver_inputs

    path = FIXTURE_DIR / "go_holdback_denver.json"
    fixture = run(path)
    inputs = denver_inputs(interest_rate="0.145")
    result = underwrite(inputs, CONFIG)
    assert result.sensitivity is not None and result.sensitivity.deal_cell is not None
    sheet = build_workbook(fixture.screen_inputs, inputs, result, CONFIG)["Lender Sensitivity"]
    bottom = 4 + len(result.sensitivity.rows)
    column = 5 + result.sensitivity.rates.index(Decimal("0.145"))
    shaded = sheet.cell(row=bottom, column=column)
    assert shaded.fill.fgColor.rgb == "00E7F3EC" and shaded.font.bold
    assert Decimal(str(shaded.value)) == result.return_overview.irr
    assert sheet.cell(row=bottom, column=column - 1).fill.fgColor.rgb != "00E7F3EC"


def test_a_result_from_before_the_table_still_exports_and_says_so() -> None:
    path = FIXTURE_DIR / "go_holdback_denver.json"
    fixture = run(path)
    assert fixture.underwrite_inputs is not None and fixture.underwrite_result is not None
    older = fixture.underwrite_result.model_copy(update={"sensitivity": None})
    sheet = build_workbook(fixture.screen_inputs, fixture.underwrite_inputs, older, CONFIG)[
        "Lender Sensitivity"
    ]
    assert "predates the sensitivity table" in str(sheet.cell(row=3, column=1).value)


# --- the borrower sensitivity and pricing sheets (SPEC §8.9, §8.10) ----------------------------


def test_the_borrower_sheet_is_profit_cost_and_project_yield_under_each_rate() -> None:
    fixture = run(FIXTURE_DIR / "go_acquisition_wichita.json")
    result = fixture.underwrite_result
    assert result is not None and result.borrower_sensitivity is not None
    grid = result.borrower_sensitivity
    sheet = book_for(FIXTURE_DIR / "go_acquisition_wichita.json")["Borrower Sensitivity"]
    assert [D(str(sheet.cell(row=4, column=2 + 3 * n).value)) for n in range(len(grid.rates))] == (
        grid.rates
    )
    assert [sheet.cell(row=5, column=c).value for c in range(1, 5)] == [
        "Loan amount",
        "Profit",
        "Total cost",
        "Project yield",
    ]
    bottom = 5 + len(grid.rows)
    assert D(str(sheet.cell(row=bottom, column=1).value)) == D("150000")
    profit, cost, project = (sheet.cell(row=bottom, column=c) for c in (2, 3, 4))
    # the first column is 12.5% (tests/test_borrower_sensitivity.py records the figures);
    # Wichita is priced at 12%, off the grid, so nothing is shaded
    assert (D(str(profit.value)), D(str(cost.value))) == (D("17650"), D("226750"))
    assert round(D(str(project.value)), 6) == D("0.077839")
    assert profit.fill.fgColor.rgb != "00E7F3EC" and project.number_format == "0.0%"


def test_the_pricing_sheet_carries_every_row_and_its_sentence() -> None:
    fixture = run(FIXTURE_DIR / "go_acquisition_wichita.json")
    result = fixture.underwrite_result
    assert result is not None and result.pricing is not None
    sheet = book_for(FIXTURE_DIR / "go_acquisition_wichita.json")["Pricing"]
    assert D(str(sheet.cell(row=2, column=2).value)) == D("0.20")
    assert sheet.cell(row=4, column=1).value == "Step"
    for offset, row in enumerate(result.pricing.rows, start=5):
        assert sheet.cell(row=offset, column=1).value == row.label
        assert D(str(sheet.cell(row=offset, column=2).value)) == row.loan_amount
        assert D(str(sheet.cell(row=offset, column=3).value)) == row.step
        assert D(str(sheet.cell(row=offset, column=4).value)) == row.interest_rate
        assert sheet.cell(row=offset, column=9).value == row.sentence
        # three columns on the end, so the ones before them did not move (engine 2.0.0)
        assert D(str(sheet.cell(row=offset, column=11).value)) == row.borrower_equity
        added = sheet.cell(row=offset, column=12).value
        assert (None if added is None else D(str(added))) == row.equity_added
        assert sheet.cell(row=offset, column=13).value == row.derisk
    assert [sheet.cell(row=4, column=c).value for c in (11, 12, 13)] == [
        "Borrower equity",
        "Equity added",
        "De-risk",
    ]
    # most equity first: the smallest loan at the top
    assert [sheet.cell(row=r, column=1).value for r in (5, 6, 7)] == ["-10%", "-5%", "Requested"]
    assert sheet.cell(row=5, column=13).value == (
        "Adding $15,000 of equity brings the rate to 16.34%."
    )


def test_the_pricing_sheet_carries_the_max_quote_and_the_rate_a_row_needs() -> None:
    """Wichita needs 16.34%; under a 16% maximum the Rate cell is empty and Rate needed set."""
    from config.config import DEFAULT_PATH, load_yaml
    from engine.underwrite import underwrite

    data = load_yaml(DEFAULT_PATH.read_text(encoding="utf-8"))
    data["pricing"]["max_quoted_rate"] = "0.16"
    capped = Config.from_dict(data)
    fixture = run(FIXTURE_DIR / "go_acquisition_wichita.json")
    assert fixture.underwrite_inputs is not None
    result = underwrite(fixture.underwrite_inputs, capped)
    sheet = build_workbook(fixture.screen_inputs, fixture.underwrite_inputs, result, capped)[
        "Pricing"
    ]
    assert sheet.cell(row=2, column=3).value == "Max quoted rate"
    assert D(str(sheet.cell(row=2, column=4).value)) == D("0.16")
    assert sheet.cell(row=4, column=10).value == "Rate needed"
    assert sheet.cell(row=5, column=4).value is None
    needed = sheet.cell(row=5, column=10)
    assert D(str(needed.value)) == D("0.1634") and needed.number_format == "0.00%"
    assert sheet.cell(row=5, column=9).value == (
        "A $135,000 loan exceeds max quoted rate (16.34% needed)."
    )


def test_a_borrower_sheet_with_no_sale_price_says_not_evaluated() -> None:
    fixture = run(FIXTURE_DIR / "conditional_no_price_no_rent.json")
    assert fixture.underwrite_inputs is not None and fixture.underwrite_result is not None
    sheet = build_workbook(
        fixture.screen_inputs, fixture.underwrite_inputs, fixture.underwrite_result, CONFIG
    )["Borrower Sensitivity"]
    assert "NOT EVALUATED" in str(sheet.cell(row=3, column=1).value)


def test_a_result_from_before_1_7_0_exports_both_sheets_saying_so() -> None:
    fixture = run(FIXTURE_DIR / "go_acquisition_wichita.json")
    assert fixture.underwrite_inputs is not None and fixture.underwrite_result is not None
    older = fixture.underwrite_result.model_copy(
        update={"borrower_sensitivity": None, "pricing": None}
    )
    book = build_workbook(fixture.screen_inputs, fixture.underwrite_inputs, older, CONFIG)
    assert "engine 1.7.0" in str(book["Borrower Sensitivity"].cell(row=3, column=1).value)
    assert "engine 1.7.0" in str(book["Pricing"].cell(row=3, column=1).value)


# --- the rest of the sheets ----------------------------------------------------------------------


def test_the_inputs_sheet_names_where_each_valuation_came_from() -> None:
    book = book_for(FIXTURE_DIR / "go_team_overrides_tulsa.json")
    sheet = book["Inputs"]
    notes = {row[0].value: row[2].value for row in sheet.iter_rows(min_col=1, max_col=3)}
    assert notes["Estimated sale price"] == "team, by hand"
    assert "As-is value" in notes and "Rental exit value" in notes  # SPEC §8.5, §8.6


def test_an_analysis_that_did_not_run_says_why_on_its_own_sheet() -> None:
    book = book_for(FIXTURE_DIR / "conditional_no_price_no_rent.json")
    status, _ = find_row(book["Rental"], "Status")
    assert status == "NOT_EVALUATED"
    status, _ = find_row(book["Take-Back"], "Status")
    assert status == "NOT_EVALUATED"
    take_back = book["Take-Back"]
    notes = {row[0].value: row[2].value for row in take_back.iter_rows(min_col=1, max_col=3)}
    assert "no projected sale price" in notes["Status"]
    assert "Toggle" not in notes
    rental_notes = {
        row[0].value: row[2].value for row in book["Rental"].iter_rows(min_col=1, max_col=3)
    }
    assert "no exit value" in rental_notes["Status"]
    flip = book["Flip"]
    status, _ = find_row(flip, "Status")
    assert status == "NOT_EVALUATED"
    flip_notes = {row[0].value: row[2].value for row in flip.iter_rows(min_col=1, max_col=3)}
    assert "no estimated sale price" in flip_notes["Status"]


def test_the_flip_sheet_reports_the_cost_stack_with_no_sale_price() -> None:
    book = book_for(FIXTURE_DIR / "conditional_no_price_no_rent.json")
    status, _ = find_row(book["Flip"], "Status")
    assert status == "NOT_EVALUATED"
    total_costs, fmt = find_row(book["Flip"], "Total deal cost")
    assert total_costs > 0 and fmt == '"$"#,##0.00'
    net_profit, _ = find_row(book["Flip"], "Net profit")
    assert net_profit is None
    equity, _ = find_row(book["Flip"], "Borrower equity")
    assert D(str(equity)) == D("29400")  # 109,400 of deal cost on an 80,000 loan


def test_the_flip_sheet_is_the_profit_line_by_line_then_the_cost_stack() -> None:
    """Sale price, broker, purchase, closing, rehab, contingency, holding, interest,
    origination, profit; then the stack to the total deal cost (SPEC §8.4)."""
    sheet = book_for(FIXTURE_DIR / "go_holdback_denver.json")["Flip"]
    labels = [row[0].value for row in sheet.iter_rows(min_col=1, max_col=1)]
    profit = labels[labels.index("PROFIT") + 1 : labels.index("Net profit") + 1]
    assert profit == [
        "Sale price",
        "Broker",
        "Purchase",
        "Closing",
        "Rehab",
        "Contingency",
        "Holding",
        "Interest",
        "Origination",
        "Net profit",
    ]
    stack = labels[labels.index("COST STACK") + 1 : labels.index("Total deal cost") + 1]
    assert stack == profit[2:-1] + ["Total deal cost"]
    rows = [(row[0].value, row[1].value) for row in sheet.iter_rows(min_col=1, max_col=2)]
    start = labels.index("PROFIT") + 1
    lines = [D(str(value)) for _, value in rows[start : start + len(profit) - 1]]
    assert lines[0] == D("330000") and all(line <= 0 for line in lines[1:])
    values = dict(rows)
    assert sum(lines, D(0)) == D(str(values["Net profit"])) == D("31810")
    assert D(str(values["Total deal cost"])) == D("278390")
    assert D(str(values["Borrower equity"])) == D("83390")


def test_the_sources_and_uses_sheet_balances_with_an_excel_formula() -> None:
    sheet = book_for(FIXTURE_DIR / "go_holdback_denver.json")["Sources & Uses"]
    values = {row[0].value: row[1].value for row in sheet.iter_rows(min_col=1, max_col=2)}
    assert D(str(values["Advance at closing"])) == D("147000")
    assert D(str(values["Rehab holdback"])) == D("48000")
    assert D(str(values["Borrower equity"])) == D("83390")
    assert D(str(values["Total sources"])) == D(str(values["Total uses"])) == D("278390")
    assert values["Sources - uses"] == "=B7-B17"
    assert sheet["A7"].value == "Total sources" and sheet["A17"].value == "Total uses"


def test_the_flags_sheet_lists_every_flag_with_its_severity_and_message() -> None:
    path = FIXTURE_DIR / "go_acquisition_wichita.json"
    result = run(path).underwrite_result
    assert result is not None
    sheet = book_for(path)["Flags"]
    rows = [
        (row[0].value, row[1].value, row[2].value)
        for row in sheet.iter_rows(min_row=4, max_col=3)
        if row[0].value
    ]
    assert len(rows) == len(result.flags)
    for (code, severity, message), flag in zip(rows, result.flags, strict=True):
        assert code == flag.code.value
        assert severity == flag.severity.value
        assert message == flag.message


def test_a_run_with_no_flags_says_so() -> None:
    fixture = run(FIXTURE_DIR / "go_hold_denver.json")
    assert fixture.underwrite_inputs is not None and fixture.underwrite_result is not None
    clean = fixture.underwrite_result.model_copy(update={"flags": []})
    sheet = build_workbook(fixture.screen_inputs, fixture.underwrite_inputs, clean, CONFIG)["Flags"]
    assert sheet.cell(row=4, column=1).value == "(none)"


def test_the_inputs_sheet_says_when_a_number_is_the_borrower_s_own() -> None:
    """SPEC §4.2, §6.1: the workbook names the borrower where that is the source."""
    book = book_for(FIXTURE_DIR / "go_web_borrower_values_okc.json")
    sheet = book["Inputs"]
    notes = {row[0].value: row[2].value for row in sheet.iter_rows(min_col=1, max_col=3)}
    assert notes["Estimated sale price"] == "borrower, on the public form"
    assert notes["Monthly rent"] == "borrower, on the public form"

    corrected = book_for(FIXTURE_DIR / "go_web_overridden_by_team_okc.json")["Inputs"]
    notes = {row[0].value: row[2].value for row in corrected.iter_rows(min_col=1, max_col=3)}
    assert notes["Estimated sale price"] == "team, by hand"
    assert notes["Monthly rent"] == "borrower, on the public form"
